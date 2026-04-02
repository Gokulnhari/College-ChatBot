"""
agents/report_agent.py
Generates Excel or PDF reports from the student dataframe.
"""
import io
import pandas as pd
from typing import List, Dict, Any


def _coerce_value(series: pd.Series, val: Any) -> Any:
    """Coerce filter value to match the dtype of the column."""
    if pd.api.types.is_integer_dtype(series.dtype):
        try:
            return int(float(str(val)))
        except (ValueError, TypeError):
            return val
    if pd.api.types.is_float_dtype(series.dtype):
        try:
            return float(str(val))
        except (ValueError, TypeError):
            return val
    return str(val)


def apply_filters(df: pd.DataFrame, filters: List[Dict]) -> pd.DataFrame:
    """Apply filters to dataframe with automatic type coercion."""
    result = df.copy()
    for f in filters:
        col = f.get("column") or f.get("field")
        op  = f.get("operator", "==")
        val = f.get("value")
        if not col or col not in result.columns:
            print(f"[report] Skipping unknown column: {col}")
            continue
        val = _coerce_value(result[col], val)
        print(f"[report] Filter: {col} {op} {val!r} (dtype={result[col].dtype})")
        if op == "==":
            result = result[result[col] == val]
        elif op == "!=":
            result = result[result[col] != val]
        elif op == ">":
            result = result[result[col] > val]
        elif op == "<":
            result = result[result[col] < val]
        elif op == ">=":
            result = result[result[col] >= val]
        elif op == "<=":
            result = result[result[col] <= val]
        elif op in ("contains", "like"):
            result = result[
                result[col].astype(str).str.contains(str(val), case=False, na=False)
            ]
        print(f"[report] Rows after filter: {len(result)}")
    return result


def select_columns_for_report(df: pd.DataFrame, report_type: str) -> pd.DataFrame:
    """Select relevant columns based on report type."""
    all_cols = list(df.columns)

    def find_cols(*keywords):
        return [c for c in all_cols if any(k in c.lower() for k in keywords)]

    id_cols         = find_cols("id")
    name_cols       = find_cols("name")
    class_cols      = find_cols("class", "section")
    attendance_cols = find_cols("attendance")
    marks_cols      = find_cols("marks", "score", "grade", "cgpa", "gpa")
    fee_cols        = find_cols("fee", "paid", "due", "amount")

    if report_type == "attendance":
        cols = id_cols + name_cols + class_cols + attendance_cols
    elif report_type == "marks":
        cols = id_cols + name_cols + class_cols + marks_cols
    elif report_type == "fee":
        cols = id_cols + name_cols + class_cols + fee_cols
    else:
        cols = all_cols

    seen, unique = set(), []
    for c in cols:
        if c not in seen and c in all_cols:
            seen.add(c)
            unique.append(c)

    return df[unique] if unique else df


def _generate_excel(report_df: pd.DataFrame) -> bytes:
    """Generate Excel file bytes."""
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        report_df.to_excel(writer, index=False, sheet_name="Report")
        worksheet = writer.sheets["Report"]

        # Style header row
        from openpyxl.styles import PatternFill, Font, Alignment
        header_fill = PatternFill(start_color="1a3a6b", end_color="1a3a6b", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)

        for col_idx, col in enumerate(report_df.columns, 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.fill      = header_fill
            cell.font      = header_font
            cell.alignment = Alignment(horizontal="center")

            # Auto-size column
            max_len = max(
                len(str(col)),
                report_df[col].astype(str).str.len().max()
                if not report_df.empty else 0
            )
            worksheet.column_dimensions[cell.column_letter].width = min(max_len + 2, 40)

    return buffer.getvalue()


def _generate_pdf(report_df: pd.DataFrame, title: str = "Report") -> bytes:
    """Generate PDF file bytes using reportlab."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.enums import TA_CENTER
    except ImportError:
        raise ImportError("reportlab not installed. Run: pip install reportlab")

    buffer = io.BytesIO()

    # Use landscape for wide tables
    pagesize = landscape(A4) if len(report_df.columns) > 5 else A4

    doc = SimpleDocTemplate(
        buffer, pagesize=pagesize,
        rightMargin=1.5*cm, leftMargin=1.5*cm,
        topMargin=2*cm,     bottomMargin=2*cm
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle", parent=styles["Title"],
        fontSize=16, spaceAfter=6, alignment=TA_CENTER
    )
    subtitle_style = ParagraphStyle(
        "ReportSubtitle", parent=styles["Normal"],
        fontSize=10, spaceAfter=12, alignment=TA_CENTER,
        textColor=colors.grey
    )

    elements = []
    elements.append(Paragraph(title, title_style))
    elements.append(Paragraph(
        f"Total Records: {len(report_df)}",
        subtitle_style
    ))
    elements.append(Spacer(1, 0.3*cm))

    # Build table data
    headers  = [c.replace("_", " ") for c in report_df.columns]
    rows     = report_df.astype(str).values.tolist()
    data     = [headers] + rows

    # Calculate column widths
    page_w   = pagesize[0] - 3*cm
    n_cols   = len(report_df.columns)
    col_w    = page_w / n_cols
    col_widths = [col_w] * n_cols

    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        # Header
        ("BACKGROUND",    (0, 0), (-1, 0),  colors.HexColor("#1a3a6b")),
        ("TEXTCOLOR",     (0, 0), (-1, 0),  colors.white),
        ("FONTNAME",      (0, 0), (-1, 0),  "Helvetica-Bold"),
        ("FONTSIZE",      (0, 0), (-1, 0),  9),
        ("ALIGN",         (0, 0), (-1, 0),  "CENTER"),
        # Data rows
        ("FONTNAME",      (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE",      (0, 1), (-1, -1), 8),
        ("ALIGN",         (0, 0), (-1, -1), "CENTER"),
        ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
        # Alternating row colors
        ("ROWBACKGROUNDS",(0, 1), (-1, -1), [colors.white, colors.HexColor("#f0f4ff")]),
        # Grid
        ("GRID",          (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
        ("ROWHEIGHT",     (0, 0), (-1, -1), 16),
        ("TOPPADDING",    (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))

    elements.append(table)
    doc.build(elements)
    return buffer.getvalue()


def generate_report(
    report_type: str,
    filters: List[Dict],
    df: pd.DataFrame,
    output_format: str = "excel"
) -> bytes:
    """
    Generate a report as Excel or PDF bytes.
    Returns file bytes ready to be sent as HTTP response.
    """
    print(f"[report] type={report_type} | format={output_format} | filters={filters}")
    print(f"[report] Input df shape: {df.shape}")

    # Apply filters
    filtered_df = apply_filters(df, filters) if filters else df.copy()
    print(f"[report] After filters: {filtered_df.shape}")

    if filtered_df.empty:
        print("[report] WARNING: No rows matched the filters!")

    # Select relevant columns
    report_df = select_columns_for_report(filtered_df, report_type)
    report_df = report_df.sort_values(report_df.columns[0]).reset_index(drop=True)
    print(f"[report] Final report shape: {report_df.shape}")

    # Build a human-readable title
    filter_parts = []
    for f in (filters or []):
        col = f.get("column") or f.get("field", "")
        val = f.get("value", "")
        filter_parts.append(f"{col} {val}")
    filter_str = " | ".join(filter_parts) if filter_parts else "All Records"
    title = f"{report_type.title()} Report — {filter_str}"

    if output_format == "pdf":
        return _generate_pdf(report_df, title=title)
    else:
        return _generate_excel(report_df)